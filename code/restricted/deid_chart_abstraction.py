#!/usr/bin/env python3
"""De-identify the completed chart abstraction workbook. RUNS LOCALLY ONLY.

Reads  Chart Abstraction.xlsx  (contains real dates = PHI)
Writes deid/chart_*.csv         (dates replaced by integer days from that patient's biopsy)

The biopsy anchor comes from PHI_CROSSWALK_DO_NOT_SHARE.csv, which never leaves this machine.
Intervals are preserved exactly; no calendar date is emitted.
"""
import openpyxl, csv, os, re, datetime

ROOT = os.path.expanduser("~/mnt/Organotypic Slice Culture")
SRC  = os.path.join(ROOT, "Chart Abstraction.xlsx")
OUT  = os.path.join(ROOT, "deid")
XW   = os.path.join(ROOT, "PHI_CROSSWALK", "PHI_CROSSWALK_DO_NOT_SHARE.csv")
os.makedirs(OUT, exist_ok=True)

# ---- biopsy anchor per patient
ANCHOR = {}
with open(XW) as fh:
    for row in csv.DictReader(fh):
        s = (row.get("Biopsy_Date") or "").strip()
        for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
            try:
                ANCHOR[row["study_id"]] = datetime.datetime.strptime(s, fmt).date(); break
            except Exception:
                pass

DATEPAT = [re.compile(r'\b\d{1,2}[/-]\d{1,2}[/-]\d{2,4}\b'),
           re.compile(r'\b\d{4}[/-]\d{1,2}[/-]\d{1,2}\b'),
           re.compile(r'\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+\d{1,2},?\s*\d{0,4}\b', re.I)]
LONGNUM = re.compile(r'\b\d{6,}\b')
MRNLIKE = re.compile(r'\b[A-Z]{0,2}\d{7,9}\b')

def as_date(v):
    if v is None: return None
    if isinstance(v, datetime.datetime): return v.date()
    if isinstance(v, datetime.date): return v
    s = str(v).strip()
    for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d", "%m-%d-%Y", "%m-%d-%y"):
        try: return datetime.datetime.strptime(s, fmt).date()
        except Exception: pass
    return None

def rel(v, sid):
    """calendar date -> integer days from that patient's biopsy"""
    d, a = as_date(v), ANCHOR.get(sid)
    if d is None or a is None: return ""
    return (d - a).days

def scrub(text, sid):
    """free text: dates -> [D+n], long numbers -> [ID]"""
    if text is None: return ""
    s = str(text)
    a = ANCHOR.get(sid)
    def repl(m):
        d = as_date(m.group(0))
        if d is None or a is None: return "[DATE]"
        return "[D%+d]" % (d - a).days
    for p in DATEPAT: s = p.sub(repl, s)
    s = MRNLIKE.sub("[ID]", s); s = LONGNUM.sub("[ID]", s)
    return " ".join(s.split())

wb = openpyxl.load_workbook(SRC, data_only=True)

def grab(sheet, date_cols, text_cols, out_name, rename=None):
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows[:12])
              if r and sum(1 for v in r if v not in (None, "")) >= 3)
    hdr = [str(v).strip() if v is not None else "" for v in rows[hi]]
    body = [r for r in rows[hi+1:] if r and any(v not in (None, "") for v in r)]
    keep, dropped = [], 0
    for r in body:
        rec = {}
        for j, h in enumerate(hdr):
            if not h: continue
            rec[h] = r[j] if j < len(r) else None
        sid = str(rec.get("Study ID") or "").strip()
        # drop the example row and anything without a study id
        if not sid.upper().startswith("PT-"):
            dropped += 1; continue
        out = {"study_id": sid}
        for h, v in rec.items():
            if h == "Study ID": continue
            key = (rename or {}).get(h, h)
            if h == "Biopsy":       out[key] = re.sub(r"\s*\([^)]*\)", "", str(v or "")).strip()
            elif h in date_cols:    out[key] = rel(v, sid)
            elif h in text_cols:    out[key] = scrub(v, sid)
            else:                   out[key] = "" if v is None else v
        keep.append(out)
    if not keep:
        print(f"  {sheet}: no usable rows"); return
    cols = list(keep[0].keys())
    path = os.path.join(OUT, out_name)
    with open(path, "w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=cols, extrasaction="ignore")
        w.writeheader(); w.writerows(keep)
    print(f"  {sheet:<14} -> {out_name:<28} {len(keep)} rows ({dropped} non-PT rows dropped)")

grab("1_Therapy",
     {"Start date", "Stop date", "Date best response assessed", "Date of progression"},
     {"Notes", "Biopsy"}, "chart_therapy_deid.csv",
     rename={"Start date":"day_start","Stop date":"day_stop",
             "Date best response assessed":"day_best_response",
             "Date of progression":"day_progression","Line #":"line",
             "RECIST best response":"best_response","Reason stopped":"reason_stopped",
             "Cycles completed":"cycles","Biopsy":"biopsy","Regimen":"regimen","Notes":"notes"})

grab("2_Markers", {"Collection date"}, {"Notes"}, "chart_markers_deid.csv",
     rename={"Collection date":"day","Test":"test","Result":"result","Units":"units",
             "Below detection? (Y/N)":"below_detection","Notes":"notes"})

grab("3_Imaging", {"Scan date"}, {"Notes"}, "chart_imaging_deid.csv",
     rename={"Scan date":"day","Modality":"modality","Target lesion sum (mm)":"target_sum_mm",
             "RECIST assessment":"recist","New lesions? (Y/N)":"new_lesions",
             "Regimen on at time of scan":"regimen_at_scan","Notes":"notes"})

grab("4_FollowUp", {"Date of last clinical contact", "Date of death"}, {"Notes","Cause of death"},
     "chart_followup_deid.csv",
     rename={"Date of last clinical contact":"day_last_contact",
             "Vital status (Alive/Deceased)":"vital_status","Date of death":"day_death",
             "Cause of death":"cause_of_death",
             "Still on treatment at last contact? (Y/N)":"on_treatment_at_last_contact",
             "Notes":"notes"})

# ERK methods sheet has no patient data and no dates - copy verbatim
ws = wb["5_ERK_methods"]
rows = [r for r in ws.iter_rows(values_only=True) if r and any(v not in (None,"") for v in r)]
with open(os.path.join(OUT, "chart_erk_methods.csv"), "w", newline="") as fh:
    csv.writer(fh).writerows([[("" if v is None else v) for v in r] for r in rows])
print(f"  5_ERK_methods  -> chart_erk_methods.csv        {len(rows)} rows")

# ---- leak check: no calendar date may survive
import glob
bad = []
for p in glob.glob(os.path.join(OUT, "chart_*.csv")):
    for i, line in enumerate(open(p), 1):
        for pat in DATEPAT:
            if pat.search(line): bad.append((os.path.basename(p), i, pat.pattern[:24]))
print("\nLEAK CHECK:", "CLEAN - no calendar dates in output" if not bad else f"FAILED {bad[:6]}")
