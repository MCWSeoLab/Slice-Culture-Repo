#!/usr/bin/env python3
"""HIPAA Safe Harbor de-identification - Organotypic Slice Culture project. Runs LOCALLY."""
import openpyxl, csv, os, re, datetime
from collections import OrderedDict

ROOT = os.path.expanduser("~/mnt/Organotypic Slice Culture")
OUT, XW = os.path.join(ROOT,"deid"), os.path.join(ROOT,"PHI_CROSSWALK")
os.makedirs(OUT, exist_ok=True); os.makedirs(XW, exist_ok=True)
F_CLINIC = os.path.join(ROOT,"Donutsuccess_clinic correlation.xlsx")
F_DATA   = os.path.join(ROOT,"Organotypic Cultures Data (UPDATED 7_2026) (3).xlsx")

audit=[]
def A(m): audit.append(str(m))

DATE_PATTERNS=[re.compile(r'\b\d{1,2}[/-]\d{1,2}[/-]\d{2,4}\b'),
               re.compile(r'\b\d{4}[/-]\d{1,2}[/-]\d{1,2}\b'),
               re.compile(r'\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+\d{1,2},?\s*\d{0,4}\b',re.I)]
LONGNUM=re.compile(r'\b\d{6,}\b'); PHONE=re.compile(r'\b\d{3}[-.\s]\d{3}[-.\s]\d{4}\b')
EMAIL=re.compile(r'\b[\w.+-]+@[\w.-]+\.\w+\b'); MCWMRN=re.compile(r'\b[A-Z]{0,2}\d{7,9}\b')

def as_date(v):
    if v is None: return None
    if isinstance(v,datetime.datetime): return v.date()
    if isinstance(v,datetime.date): return v
    if isinstance(v,(int,float)):
        try:
            f=float(v)
            if 20000 < f < 60000: return datetime.date(1899,12,30)+datetime.timedelta(days=f)
        except Exception: pass
        return None
    s=str(v).strip()
    for fmt in ("%m/%d/%Y","%m/%d/%y","%Y-%m-%d","%m-%d-%Y","%m-%d-%y"):
        try: return datetime.datetime.strptime(s,fmt).date()
        except Exception: pass
    return None

STOP={"no","yes","na","nan","none","pr","sd","pd","cr","ne","gem","nab","tx","ca","ih",
      "er","or","and","the","of","on","in","at","to","dx","bx","pt","mg","ml","um","nm",
      "xt","xr","pa","br","tb","id","ok","up","hx","po","iv","d1","d2","cy","wk","mo"}
def _usable(tok):
    if not tok: return False
    t=tok.strip()
    if t.lower() in STOP: return False
    if t.isdigit(): return len(t)>=3
    return len(t)>=3

def scrub(text, anchor=None, name_tokens=()):
    if text is None: return ""
    s=str(text)
    def repl(m):
        d=as_date(m.group(0))
        if d is None or anchor is None: return "[DATE]"
        return "[D%+d]" % (d-anchor).days
    s=EMAIL.sub("[EMAIL]",s); s=PHONE.sub("[PHONE]",s)
    for p in DATE_PATTERNS: s=p.sub(repl,s)
    s=MCWMRN.sub("[ID]",s); s=LONGNUM.sub("[ID]",s)
    for tok in name_tokens:
        if _usable(tok):
            s=re.sub(r'\b'+re.escape(tok)+r'\b', "[ID]" if tok.isdigit() else "[NAME]", s, flags=re.I)
    return " ".join(s.split())

def rel(d,a): return "" if (d is None or a is None) else (d-a).days
def txt(v): return "" if v is None else " ".join(str(v).split())

def fill_merges(ws):
    grid={}
    for row in ws.iter_rows():
        for c in row: grid[(c.row,c.column)]=c.value
    for rng in ws.merged_cells.ranges:
        v=grid.get((rng.min_row,rng.min_col))
        for rr in range(rng.min_row,rng.max_row+1):
            for cc in range(rng.min_col,rng.max_col+1): grid[(rr,cc)]=v
    return grid

registry=OrderedDict()
def key_of(mrn,tb):
    k=re.sub(r'\D','',str(mrn)) if mrn is not None else ""
    k=k.lstrip("0")
    if k: return "M:"+k
    if tb is not None and txt(tb): return "T:"+re.sub(r'\W','',str(tb)).upper()
    return None
def get_id(mrn,tb,ini):
    k=key_of(mrn,tb)
    if k is None: return None
    if k not in registry:
        registry[k]={"study_id":"PT-%03d"%(len(registry)+1),"mrn":txt(mrn),"tb":txt(tb),
                     "initials":txt(ini),"name":"","biopsy_date":None}
    else:
        r=registry[k]
        if not r["tb"] and tb is not None: r["tb"]=txt(tb)
        if not r["initials"] and ini is not None: r["initials"]=txt(ini)
        if not r["mrn"] and mrn is not None: r["mrn"]=txt(mrn)
    return registry[k]["study_id"]
def rec_for(sid):
    for v in registry.values():
        if v["study_id"]==sid: return v

C={"tb":1,"mrn":2,"initials":3,"donut":4,"dx_date":5,"pre_ca199":6,"pre_reg":7,"pre_start":8,
   "pre_end":9,"pre_cycles":10,"pre_resp":11,"pre_best":12,"bx_date":13,"tested":14,
   "lab_resp":15,"mts24":16,"mts48":17,"delta":18,"tempus_pa":19,"tempus_br":20,
   "tempus_tmb":21,"tempus_xr":22,"ihc":23,"post_ca199":24,"post_reg":25,"post_start":26,
   "post_end":27,"post_cycles":28,"post_resp":29,"post_best":30,"status":31,"pre_match":33,
   "post_match":34,"which_matched":35,"lab_resp2":36,"clinic_resp":37,"correlation":38,
   "chart_date":39,"notes":40}

LABEL={"Data 3.0":"master","S + received tx":"sens_and_received_tx",
       "No Response AT ALL":"no_exvivo_response",
       "S  did NOT received tested tx A":"sens_did_not_receive_tx",
       "ViabilityNo SurvivalNo tested t":"viability_only_or_no_survival"}

wb=openpyxl.load_workbook(F_DATA,data_only=True)
SHEETS=[n for n in wb.sheetnames if n!="Calculations"]
GRIDS={n:(fill_merges(wb[n]),wb[n].max_row) for n in SHEETS}

# pass 1: register + anchor biopsy date
for n in SHEETS:
    g,mx=GRIDS[n]
    for r in range(3,mx+1):
        mrn,tb,ini=g.get((r,C["mrn"])),g.get((r,C["tb"])),g.get((r,C["initials"]))
        if not (mrn or tb): continue
        sid=get_id(mrn,tb,ini)
        if sid is None: continue
        bd=as_date(g.get((r,C["bx_date"])))
        if bd and rec_for(sid)["biopsy_date"] is None: rec_for(sid)["biopsy_date"]=bd

# pass 2: emit
pt_map, biopsies, regimens = OrderedDict(), [], []
seen_bx, seen_reg = set(), set()
for n in SHEETS:
    g,mx=GRIDS[n]; SL=LABEL.get(n,n)
    G=lambda r,k: g.get((r,C[k]))
    for r in range(3,mx+1):
        mrn,tb,ini=G(r,"mrn"),G(r,"tb"),G(r,"initials")
        if not (mrn or tb): continue
        sid=get_id(mrn,tb,ini)
        if sid is None: continue
        rec=rec_for(sid); anc=rec["biopsy_date"]
        nt=set()
        for src in (txt(mrn),txt(tb),txt(ini)):
            nt.add(src)
            for p in re.split(r'\W+',src): nt.add(p)
        nt=sorted({t for t in nt if len(t)>=2}, key=len, reverse=True)
        S=lambda k: scrub(G(r,k),anc,nt)          # free text: full scrub
        Q=lambda k: scrub(G(r,k),anc,())           # categorical: dates only

        row={"study_id":sid,"source_sheets":SL,
             "days_dx_to_biopsy":rel(as_date(G(r,"dx_date")),anc),
             "donut_analysis":S("donut"),"patient_status":Q("status"),
             "pre_bx_tx_matched_tested_drug":txt(G(r,"pre_match")),
             "post_bx_tx_matched_tested_drug":txt(G(r,"post_match")),
             "which_tx_matched":S("which_matched"),"lab_response":Q("lab_resp2"),
             "clinic_response":Q("clinic_resp"),"correlation_present":Q("correlation"),
             "days_biopsy_to_chart_review":rel(as_date(G(r,"chart_date")),anc),
             "notes":S("notes")}
        if sid not in pt_map: pt_map[sid]=row
        else:
            cur=pt_map[sid]
            for k,v in row.items():
                if k in ("study_id","source_sheets"): continue
                if cur.get(k) in (None,"") and v not in (None,""): cur[k]=v
            if SL not in cur["source_sheets"].split("|"): cur["source_sheets"]+="|"+SL

        bk=(sid,S("tested"),Q("mts24"),Q("mts48"),Q("lab_resp"))
        if any(G(r,k) is not None for k in ("tested","lab_resp","mts24","mts48","delta")) and bk not in seen_bx:
            seen_bx.add(bk)
            biopsies.append({"study_id":sid,"src_sheet":SL,"src_row":r,
                "tested_ex_vivo":S("tested"),"lab_response":Q("lab_resp"),
                "mts_24h":Q("mts24"),"mts_48h":Q("mts48"),"delta":Q("delta"),
                "tempus_xT_PA":S("tempus_pa"),"tempus_xT_BR":S("tempus_br"),
                "tempus_xT_TMB":S("tempus_tmb"),"tempus_xR":S("tempus_xr"),"ihc":S("ihc")})

        for ph,ks in (("pre",("pre_reg","pre_start","pre_end","pre_cycles","pre_resp","pre_best","pre_ca199")),
                      ("post",("post_reg","post_start","post_end","post_cycles","post_resp","post_best","post_ca199"))):
            reg,st,en,cy,rs,bs,ca=(G(r,k) for k in ks)
            if reg is None and st is None and rs is None: continue
            rr={"study_id":sid,"phase":ph,"src_sheet":SL,"src_row":r,
                "regimen":scrub(reg,anc,nt),
                "day_start_rel_biopsy":rel(as_date(st),anc),
                "day_end_rel_biopsy":rel(as_date(en),anc),
                "cycles_completed":scrub(cy,anc,()),"response":scrub(rs,anc,()),
                "best_response":scrub(bs,anc,()),"ca19_9":scrub(ca,anc,())}
            k2=(sid,ph,rr["regimen"],rr["day_start_rel_biopsy"],rr["response"],rr["best_response"])
            if k2 in seen_reg: continue
            seen_reg.add(k2); regimens.append(rr)

patients=list(pt_map.values())
A("Workbook 2: merged %d sheets (%s)" % (len(SHEETS), ", ".join(SHEETS)))
A("  patients=%d  biopsy-records=%d  regimen-records=%d"%(len(patients),len(biopsies),len(regimens)))

# ---- clinic sheet
clinic=[]
ws1=openpyxl.load_workbook(F_CLINIC,data_only=True)["Sheet1"]
rows1=list(ws1.iter_rows(min_row=1,values_only=True))
HDR=[txt(v) for v in rows1[0]]
DROP={"Patient Name","MRN","Tissue bank#","Date of Biopsy"}
keep=[i for i,h in enumerate(HDR) if h and h not in DROP]
for r in rows1[1:]:
    if not any(v is not None for v in r): continue
    name,mrn,tb=txt(r[0]),r[1],r[2]
    sid=get_id(mrn,tb,name)
    if sid is None: continue
    rec=rec_for(sid)
    if not rec["name"]: rec["name"]=name
    bd=as_date(r[3])
    if bd and rec["biopsy_date"] is None: rec["biopsy_date"]=bd
    nt=sorted({p for p in [name]+re.split(r'\W+',name)+[txt(mrn),txt(tb)] if _usable(p)},key=len,reverse=True)
    o={"study_id":sid}
    for i in keep: o[HDR[i]]=scrub(r[i] if i<len(r) else None, rec["biopsy_date"], nt)
    clinic.append(o)
A("Workbook 1 'Donutsuccess': %d patient rows"%len(clinic))

# ---- global identifier sweep
GT=set()
for v in registry.values():
    for f in ("mrn","tb","initials","name"):
        val=v.get(f) or ""
        GT.add(val)
        for p in re.split(r'\W+',val): GT.add(p)
GT=sorted({t for t in GT if _usable(t)},key=len,reverse=True)
GRX=[(t,re.compile(r'\b'+re.escape(t)+r'\b',re.I)) for t in GT]
sub=0
for coll in (patients,biopsies,regimens,clinic):
    for row in coll:
        for k,v in list(row.items()):
            if isinstance(v,str) and v:
                nv=v
                for t,rx in GRX:
                    n2=rx.sub("[ID]" if t.isdigit() else "[NAME]",nv)
                    if n2!=nv: sub+=1
                    nv=n2
                row[k]=nv
A("  global identifier sweep: %d substitutions across %d tokens"%(sub,len(GT)))

def write(p,rows,fields=None):
    if not rows: A("  (empty) "+os.path.basename(p)); return
    fs=fields or list(rows[0].keys())
    with open(p,"w",newline="",encoding="utf-8") as f:
        w=csv.DictWriter(f,fieldnames=fs,extrasaction="ignore"); w.writeheader(); w.writerows(rows)
    A("  wrote %-30s %d rows x %d cols"%(os.path.basename(p),len(rows),len(fs)))

A("OUTPUT:")
write(os.path.join(OUT,"patients_deid.csv"),patients)
write(os.path.join(OUT,"biopsy_exvivo_deid.csv"),biopsies)
write(os.path.join(OUT,"regimens_deid.csv"),regimens)
ck=[]
for c in clinic:
    for k in c:
        if k not in ck: ck.append(k)
write(os.path.join(OUT,"clinic_correlation_deid.csv"),clinic,ck)

with open(os.path.join(XW,"PHI_CROSSWALK_DO_NOT_SHARE.csv"),"w",newline="",encoding="utf-8") as f:
    w=csv.writer(f); w.writerow(["study_id","MRN","TB_number","Patient_Initials","Patient_Name","Biopsy_Date"])
    for v in registry.values():
        w.writerow([v["study_id"],v["mrn"],v["tb"],v["initials"],v["name"],
                    v["biopsy_date"].isoformat() if v["biopsy_date"] else ""])
A("  wrote crosswalk: %d patients"%len(registry))
A("  patients with no anchor biopsy date: %d"%sum(1 for v in registry.values() if v["biopsy_date"] is None))
open(os.path.join(OUT,"_AUDIT.txt"),"w").write("\n".join(audit))
print("\n".join(audit))
